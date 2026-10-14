"""Build DOCUMENTATION.docx from DOCUMENTATION.md, keeping the original file's styles.

Usage: python md2docx.py <plugin dir> [out.docx ...]
Styles come from base-original.docx ( a copy of the python-docx-built original ) beside this script.
"""
import os
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = sys.argv[1]
MD = os.path.join(ROOT, 'DOCUMENTATION.md')
BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'base-original.docx')
OUTS = sys.argv[2:] or [os.path.join(ROOT, 'DOCUMENTATION.docx'), os.path.join(ROOT, 'submission', 'DOCUMENTATION.docx')]
TEXT_W = 6.5  # inches between the 1" margins of US Letter.
MAX_H = 4.3   # tallest an image may be, in inches.

doc = docx.Document(BASE)
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

doc.core_properties.title = 'My Account Dashboard Builder — Documentation'
doc.core_properties.subject = 'How to customize the WooCommerce My Account page'

GRAY = RGBColor(0x55, 0x5F, 0x6D)

# ---------------------------------------------------------------- numbering
numbering = doc.part.numbering_part.numbering_definitions._numbering
style_num = doc.styles['List Number'].element.pPr.numPr.numId.val
abstract_id = [n for n in numbering.num_lst if n.numId == style_num][0].abstractNumId.val


def set_num(par, num_id):
    numpr = par._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = num_id


def new_list():
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


# ---------------------------------------------------------------- inline
TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))')


def add_link(par, text, url):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    par._p.append(h)


def run(par, text, bold, italic, size, color):
    r = par.add_run(text)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if size:
        r.font.size = size
    if color:
        r.font.color.rgb = color
    return r


def add_inline(par, text, bold=False, italic=False, size=None, color=None):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            run(par, text[pos:m.start()], bold, italic, size, color)
        tok = m.group(0)
        if tok.startswith('**'):
            add_inline(par, tok[2:-2], True, italic, size, color)
        elif tok.startswith('`'):
            r = run(par, tok[1:-1], bold, italic, size, color)
            r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
            r.font.size = Pt(9.5)
        elif tok.startswith('['):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok).groups()
            add_link(par, label, url)
        else:
            add_inline(par, tok[1:-1], bold, True, size, color)
        pos = m.end()
    if pos < len(text):
        run(par, text[pos:], bold, italic, size, color)


# ---------------------------------------------------------------- blocks
def picture_par():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    return p


PLACED = []          # ( section, file, caption ) for the placement table.
SECTION = ['', '']  # current H2, H3.


def kind(path):
    return 'GIF' if path.lower().endswith('.gif') else 'Image'


def file_label(par, paths):
    """A small line under the caption naming the file(s) that go here."""
    par.add_run().add_break(WD_BREAK.LINE)
    for k, path in enumerate(paths):
        if k:
            run(par, '   ·   ', False, False, Pt(8), GRAY)
        run(par, kind(path) + ': ', True, False, Pt(8), GRAY)
        r = run(par, path, False, False, Pt(8), GRAY)
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')


def caption(text, paths):
    c = doc.add_paragraph(style='Caption')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(12)
    add_inline(c, text, italic=True, color=GRAY)
    file_label(c, paths)
    where = SECTION[0] + ( ' › ' + SECTION[1] if SECTION[1] else '' )
    for path in paths:
        PLACED.append((where, path, text))


def add_image(path, text, width=None):
    full = os.path.join(ROOT, path)
    with Image.open(full) as im:
        px, py = im.size
    w = float(width) if width else min(TEXT_W, px / 120.0)
    # Keep tall shots to part of a page so text flows around them.
    w = min(w, MAX_H * px / py)
    picture_par().add_run().add_picture(full, width=Inches(w))
    caption(text, [path])


def add_pair(m1, m2):
    p = picture_par()
    for k, m in enumerate((m1, m2)):
        if k:
            p.add_run('  ')
        p.add_run().add_picture(os.path.join(ROOT, m.group(2)), width=Inches(float(m.group(3))))
    caption('Left: ' + m1.group(1) + ' Right: ' + m2.group(1), [m1.group(2), m2.group(2)])
    # The table lists each file with its own caption.
    PLACED[-2] = (PLACED[-2][0], m1.group(2), m1.group(1) + ' (left)')
    PLACED[-1] = (PLACED[-1][0], m2.group(2), m2.group(1) + ' (right)')


def add_code(code):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.15)
    pf.space_before = Pt(4)
    pf.space_after = Pt(10)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F3F4F6')
    p._p.get_or_add_pPr().append(shd)
    for k, line in enumerate(code):
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
        r.font.size = Pt(9)
        if k < len(code) - 1:
            r.add_break(WD_BREAK.LINE)


IMG = re.compile(r'^!\[(.*)\]\((\S+?)(?:\s+"width=([\d.]+)")?\)$')
BULLET = re.compile(r'^(\s*)[-*] (.*)$')
NUMBER = re.compile(r'^(\s*)\d+\. (.*)$')

lines = open(MD, encoding='utf-8').read().replace('\r\n', '\n').split('\n')
i = 0
para = []
num_id = None
seen_h2 = False


def flush():
    global para
    if para:
        text = ' '.join(x.strip() for x in para)
        p = doc.add_paragraph()
        if text.startswith('> '):
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, text[2:], italic=True, color=GRAY)
        else:
            add_inline(p, text)
        para = []


while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped.startswith('```'):
        flush()
        i += 1
        code = []
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code.append(lines[i])
            i += 1
        add_code(code)
        i += 1
        continue

    if not stripped:
        flush()
        i += 1
        continue

    if stripped == '---':
        flush()
        num_id = None
        i += 1
        continue

    h = re.match(r'^(#{1,4}) (.*)$', line)
    if h:
        flush()
        num_id = None
        level = len(h.group(1))
        plain = re.sub(r'[*`]', '', h.group(2)).strip()
        if level == 2:
            SECTION[0], SECTION[1] = plain, ''
        elif level == 3:
            SECTION[1] = plain
        add_inline(doc.add_heading(level=level), h.group(2))
        i += 1
        continue

    m = IMG.match(stripped)
    if m:
        flush()
        # Two narrow images in a row ( phone shots ) sit side by side.
        j = i + 1
        while j < len(lines) and not lines[j].strip():
            j += 1
        m2 = IMG.match(lines[j].strip()) if j < len(lines) else None
        if m2 and m.group(3) and m2.group(3) and float(m.group(3)) <= 3 and float(m2.group(3)) <= 3:
            add_pair(m, m2)
            i = j + 1
            continue
        add_image(m.group(2), m.group(1), m.group(3))
        i += 1
        continue

    b = BULLET.match(line)
    n = NUMBER.match(line)
    if b or n:
        flush()
        indent = len((b or n).group(1))
        text = (b or n).group(2)
        if b:
            p = doc.add_paragraph(style='List Bullet 2' if indent >= 2 else 'List Bullet')
        else:
            p = doc.add_paragraph(style='List Number')
            if num_id is None:
                num_id = new_list()
            set_num(p, num_id)
        # Continuation lines: indented, no marker.
        j = i + 1
        while j < len(lines) and lines[j].strip() and lines[j].startswith(' ') \
                and not BULLET.match(lines[j]) and not NUMBER.match(lines[j]) and not IMG.match(lines[j].strip()):
            text += ' ' + lines[j].strip()
            j += 1
        add_inline(p, text)
        i = j
        continue

    # A plain paragraph ends any numbered list.
    num_id = None
    para.append(line)
    i += 1

flush()


# ---------------------------------------------------------------- placement table
def cell_text(cell, text, bold=False, mono=False, size=Pt(8.5)):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = run(p, text, bold, False, size, None)
    if mono:
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')


def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)


doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
add_inline(doc.add_heading(level=2), 'Image and GIF Placement')
add_inline(doc.add_paragraph(),
           'Every screenshot and GIF in this guide, in order, with the section it belongs to. The files are in '
           '`submission/docs-media/` in the plugin repository; each one is also named under its caption above. '
           'GIFs animate in Word and in browsers; a PDF shows their first frame.')

COLS = [Inches(0.35), Inches(1.9), Inches(2.05), Inches(2.2)]
table = doc.add_table(rows=1, cols=4)
table.style = doc.styles['Table Grid']
table.autofit = False
for k, title in enumerate(['#', 'Where it goes', 'File', 'What it shows']):
    c = table.rows[0].cells[k]
    cell_text(c, title, bold=True)
    shade(c, 'E8EEFB')
for n, (where, path, text) in enumerate(PLACED, 1):
    row = table.add_row().cells
    cell_text(row[0], str(n))
    cell_text(row[1], where)
    cell_text(row[2], kind(path) + ': ' + os.path.basename(path), mono=True, size=Pt(8))
    cell_text(row[3], text)
    if path.lower().endswith('.gif'):
        for c in row:
            shade(c, 'FFF7E6')
for r in table.rows:
    for k, w in enumerate(COLS):
        r.cells[k].width = w
# Repeat the header row on every page.
trpr = table.rows[0]._tr.get_or_add_trPr()
hdr = OxmlElement('w:tblHeader')
hdr.set(qn('w:val'), 'true')
trpr.append(hdr)
add_inline(doc.add_paragraph(), '*GIF rows are shaded.*', color=GRAY)

for out in OUTS:
    doc.save(out)
    print('saved', out, os.path.getsize(out) // 1024, 'KB')
