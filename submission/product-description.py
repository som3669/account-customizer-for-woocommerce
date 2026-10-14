"""Build PRODUCT-DESCRIPTION.docx: the WooCommerce.com product description, laid out like the
Marketplace page ( tagline, Highlights, pitch, feature cards, Features, feature sections, Requirements ).

Usage ( from the plugin folder ): python submission/product-description.py
Images come from submission/docs-media/; styles from submission/base-original.docx.
"""
import os

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
MEDIA = os.path.join(HERE, 'docs-media')
OUT = os.path.join(HERE, 'PRODUCT-DESCRIPTION.docx')

INK = RGBColor(0x1D, 0x23, 0x27)
GRAY = RGBColor(0x55, 0x5F, 0x6D)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)

# ------------------------------------------------------------------ copy
TITLE = 'My Account Dashboard Builder'
TAGLINE = ('Branded My Account dashboard: drag-and-drop menu, Design Studio, Buy again, returns, '
           'personal offers and self-service tools.')

HIGHLIGHTS = [
    'A drag-and-drop account menu with your own pages, groups and links',
    'A Design Studio with a live preview and ready-made looks',
    'Personal coupon offers and banners for the right customers',
    'Self-service returns, cancellations and an address book',
    'View My Account exactly as any customer sees it',
]

PITCH_TITLE = 'Make My Account Work as Hard as Your Shop Pages'
PITCH = [
    'Every logged-in customer lands on My Account, yet in most stores it is a plain list of links. '
    'It is where customers check an order, change an address and decide whether to buy again, and '
    'where they get stuck and email you instead.',
    'My Account Dashboard Builder turns it into a dashboard that matches your brand: a menu customers '
    'find their way around, account numbers and order tracking at a glance, offers that bring them '
    'back, and tools that answer their own questions. No theme templates to edit and no code to write.',
]

CARDS = [
    ('menu', 'Drag-and-Drop Menu', 'Reorder, rename, hide and nest account pages on a canvas drawn the way customers see it.'),
    ('endpoints', 'Your Own Account Pages', 'Add endpoints, links, WordPress pages and collapsible groups, with Classic or Block editor content.'),
    ('rules', 'The Right Items for the Right People', 'Show an item by role, dates, products bought, orders placed, amount spent or time since the last order.'),
    ('groups', 'Menus for Customer Groups', 'Give wholesale buyers or members a menu of their own, while item settings stay shared.'),
    ('design', 'Design Studio', 'Change colours, layout, density and type beside a live preview at desktop, tablet and phone widths.'),
    ('arrange', 'Arrange the Dashboard', 'Put account numbers, the profile meter, shortcut tiles, order tracking and Buy again in the order you want.'),
    ('viewas', 'View as a Customer', 'See any customer’s menu, badges and offers, plus what is hidden from them and why. Nothing changes for them.'),
    ('offers', 'Personal Offers', 'Banners that hand each customer a one-use coupon of their own, with Copy and Use it now.'),
    ('returns', 'Returns and Cancellations', 'Customers cancel unshipped orders and request returns with a photo; you answer from one screen.'),
    ('buyagain', 'Buy Again', 'Past purchases one click from the cart, plus recently viewed products and order tracking.'),
    ('fields', 'Customer Fields', 'Ask for a VAT number, a birthday or “How did you hear about us?” at registration and in account details.'),
    ('insights', 'Insights', 'See which account pages customers open, which banners they click and the sales personal offers bring.'),
]

FEATURES = [
    'Drag-and-drop menu canvas, with keyboard moves',
    'Custom endpoints, links, WordPress pages and groups',
    'Font Awesome icons or your own images',
    'Badges and live status badges (“1 to pay”, “1 open”)',
    'Visibility rules by role, date, purchases and order history',
    'Separate menus for customer groups',
    'Design Studio with six ready-made looks, plus your own',
    'Account numbers, profile meter and shortcut tiles on the dashboard',
    'Banners with personal coupon offers',
    'Returns, order cancelling, address book and privacy page',
    'Buy again, recently viewed and order tracking',
    'Custom registration and account fields',
    'Smart tags such as {first_name} and {total_spent}',
    'Insights, import / export, WPML and Polylang ready',
]

# Sections: ( level, heading, paragraphs, bullets, media ). level 2 = heading, 0 = bold label.
SECTIONS = [
    (2, 'Build the Menu on a Canvas',
     ['The Menu Items screen draws your account menu exactly as customers see it. Drag an item to '
      'move it or drop it into a group, flip its switch to hide it, and click it to edit its label, '
      'icon and badge. The canvas follows as you type.'],
     [], ['05-drag-menu-items.gif']),
    (2, 'Add Your Own Account Pages',
     ['Click Add to menu, pick a type and type a label. Fill a page with the Classic or Block editor, '
      'shown before, after or instead of WooCommerce’s own content. Smart tags greet each customer by name.'],
     [('Endpoint', 'a page of its own under My Account, e.g. /my-account/help-center/'),
      ('Group', 'a collapsible heading that holds other items'),
      ('Link', 'any URL, in the same or a new tab'),
      ('Page', 'one of your WordPress pages')],
     ['12-front-group-and-endpoint.png']),
    (2, 'Show Each Item to the Right Customers',
     ['Every item and banner has the same visibility rules: user roles, a date range, products bought, '
      'at least or at most a number of orders, an amount spent, or a last order more than N days ago. '
      'A VIP page for big spenders or a “we miss you” offer takes a minute to set up.'],
     [], ['08-item-visibility.png']),
    (0, 'A Menu for Every Customer Group',
     ['Wholesale buyers or members can get a menu of their own. It starts as a copy of the main menu; '
      'take items out, add others and reorder it. Everyone else keeps the main menu.'],
     [], ['44-wholesale-customer.png']),
    (2, 'Design It Beside a Live Preview',
     ['The Design Studio shows the real account page at desktop, tablet and phone widths while you '
      'change colours, typeface, density, the menu style and the profile card. Start from a look:'],
     [('Classic Sidebar', 'a left menu with a subtle active bar'),
      ('Modern Cards', 'raised cards with icon chips'),
      ('Rounded Pills', 'soft pill items with a coloured active state'),
      ('Tabbed Top', 'a tab bar above the content'),
      ('Minimal', 'flat and borderless'),
      ('Theme Native', 'your theme’s own styling')],
     ['14-design-looks.gif']),
    (0, 'Arrange the Dashboard',
     ['Drag the greeting, account numbers, profile meter, shortcut tiles, order tracking and Buy again '
      'into any order, or hide them. The preview follows at once, and nothing goes live until you save.'],
     [], ['18-arrange-dashboard.gif']),
    (2, 'See My Account as Any Customer',
     ['Pick a customer in View as and the preview shows their menu, badges, offers and dashboard. A bar '
      'lists everything hidden from them, and why. It only looks: nothing is changed, sent or counted for them.'],
     [], ['21-view-as-customer.gif']),
    (2, 'Banners and Personal Offers',
     ['Put banners at the top or bottom of any account page, with an icon, text, a link and an order or '
      'cart count. Switch on a personal offer and each customer who sees the banner gets a one-use coupon '
      'for their own email, valid for the days you choose.'],
     [], ['26-front-offer-banner.png']),
    (2, 'Self-Service Returns and Cancellations',
     ['Customers cancel an order that has not shipped, within the hours you allow, and request a return '
      'on a completed order with a reason and a photo. You answer each request on the Returns tab, and '
      'the customer is emailed.'],
     [], ['29-front-order-actions.png']),
    (0, 'Buy Again, Recently Viewed and Order Tracking',
     ['Past purchases sit one click from the cart, recently viewed products wait on their own page, and '
      'recent orders show their progress on the dashboard.'],
     [], ['38-buy-again.png']),
    (2, 'Ask Customers the Questions You Need',
     ['Add a VAT number, a company, a birthday or a dropdown to the registration form and account details. '
      'Answers show on the admin order screen and work as smart tags.'],
     [], ['28-front-account-fields.png']),
    (2, 'Know What Gets Used',
     ['Insights counts account page views, banner clicks and the sales personal offers bring, over 7, 30 '
      'or 90 days. Only daily totals are kept.'],
     [], ['33-insights.png']),
    (0, 'Works on Phones',
     ['On small screens the menu folds into an Account menu button.'],
     [], ['43-front-phone.png', '43b-front-phone-menu.png']),
]

REQUIREMENTS = [
    'WordPress 6.0+, WooCommerce 6.0+ and PHP 7.4+.',
    'Compatible with HPOS and the Cart and Checkout blocks. Translation ready, with WPML and Polylang support.',
]

# Image widths, in inches ( the rest fit 6.5" wide, at most 4.3" tall ).
WIDTHS = {'05-drag-menu-items.gif': 2.9, '26-front-offer-banner.png': 4.6, '29-front-order-actions.png': 4.6,
          '28-front-account-fields.png': 3.8, '43-front-phone.png': 2.2, '43b-front-phone-menu.png': 2.2}

# ------------------------------------------------------------------ helpers
doc = docx.Document(os.path.join(HERE, 'base-original.docx'))
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
doc.core_properties.title = TITLE + ' — Product description'


def shade(el, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    el.append(shd)


def borders(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        spec = edges.get(edge)
        if spec:
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(spec[0]))
            e.set(qn('w:color'), spec[1])
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tcpr.append(b)


def para(text='', bold=False, size=None, color=None, align=None, after=8, style=None, container=None):
    p = (container or doc).add_paragraph(style=style) if style else (container or doc).add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def label(p, files):
    """Tiny gray line naming the file(s) to upload here."""
    r = p.add_run()
    r.add_break(WD_BREAK.LINE)
    for k, f in enumerate(files):
        kind = 'GIF' if f.endswith('.gif') else 'Image'
        t = p.add_run(('   ·   ' if k else '') + kind + ': ' + f)
        t.font.size = Pt(7.5)
        t.font.color.rgb = GRAY
        t.font.name = 'Consolas'
        t._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')


def media(files):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    for k, f in enumerate(files):
        path = os.path.join(MEDIA, f)
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        width = WIDTHS.get(f, min(6.5, w / 120.0))
        width = min(width, 4.3 * w / h)
        if k:
            p.add_run('  ')
        p.add_run().add_picture(path, width=Inches(width))
    label(p, files)


def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), 'D9DEE7')
    bdr.append(bottom)
    ppr.append(bdr)


# ------------------------------------------------------------------ build
h = doc.add_heading(TITLE, level=1)
para(TAGLINE, size=12, color=GRAY, after=12)

# Highlights box.
t = doc.add_table(rows=1, cols=1)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = t.rows[0].cells[0]
cell.width = Inches(6.5)
shade(cell._tc.get_or_add_tcPr(), 'F1F5FE')
borders(cell, left=(24, '1D4ED8'))
p = cell.paragraphs[0]
r = p.add_run('Highlights')
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = ACCENT
p.paragraph_format.space_after = Pt(4)
for item in HIGHLIGHTS:
    bp = cell.add_paragraph(style='List Bullet')
    br = bp.add_run(item)
    br.bold = True
    bp.paragraph_format.space_after = Pt(2)
para(after=6)

# Pitch.
doc.add_heading(PITCH_TITLE, level=2)
for text in PITCH:
    last = para(text, after=8)
last.paragraph_format.keep_with_next = True
media(['01-my-account-dashboard.png'])

# Feature cards: icon, title, one sentence; a hairline between cards.
cards = doc.add_table(rows=len(CARDS), cols=2)
cards.alignment = WD_TABLE_ALIGNMENT.CENTER
cards.autofit = False
for row, (icon, title, text) in zip(cards.rows, CARDS):
    ic, tx = row.cells
    ic.width, tx.width = Inches(0.8), Inches(5.7)
    for c in (ic, tx):
        borders(c, bottom=(4, 'E3E7EF'))
    ip = ic.paragraphs[0]
    ip.paragraph_format.space_before = Pt(6)
    ip.paragraph_format.space_after = Pt(6)
    ip.add_run().add_picture(os.path.join(MEDIA, 'icons', 'icon-' + icon + '.png'), width=Inches(0.5))
    tp = tx.paragraphs[0]
    tp.paragraph_format.space_before = Pt(6)
    tp.paragraph_format.space_after = Pt(1)
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(12)
    tr.font.color.rgb = INK
    dp = tx.add_paragraph()
    dp.paragraph_format.space_after = Pt(6)
    dp.add_run(text)
    label(dp, ['icons/icon-' + icon + '.png'])
para(after=6)

# Features list.
para('Features:', bold=True, size=12, after=4)
for item in FEATURES:
    para(item, style='List Bullet', after=1)
para(after=6)

# Feature sections: heading, text and image stay on one page.
for level, title, paras, bullets, files in SECTIONS:
    hr()
    block = []
    if level:
        block.append(doc.add_heading(title, level=level))
    else:
        block.append(para(title, bold=True, size=12, after=4))
    for text in paras:
        block.append(para(text, after=6))
    for name, text in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1)
        bp.add_run(name).bold = True
        bp.add_run(' — ' + text)
        block.append(bp)
    if bullets:
        block.append(para(after=2))
    for b in block:
        b.paragraph_format.keep_with_next = True
    media(files)

hr()
doc.add_heading('Requirements', level=3)
for text in REQUIREMENTS:
    para(text, after=2)

# Upload checklist.
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading('Media to Upload', level=2)
para('Files are in submission/docs-media/ ( icons in its icons/ folder ). The product icon and gallery '
     'images are listed in submission/media/README.md.', color=GRAY, after=8)
rows = [('Under the pitch', '01-my-account-dashboard.png')]
rows += [('Card: ' + title, 'icons/icon-' + icon + '.png') for icon, title, _ in CARDS]
rows += [(title, ', '.join(files)) for _, title, _, _, files in SECTIONS]
tbl = doc.add_table(rows=1, cols=2)
tbl.style = doc.styles['Table Grid']
for k, head in enumerate(('Where', 'File')):
    c = tbl.rows[0].cells[k]
    c.paragraphs[0].add_run(head).bold = True
    shade(c._tc.get_or_add_tcPr(), 'E8EEFB')
for where, f in rows:
    a, b = tbl.add_row().cells
    a.paragraphs[0].add_run(where).font.size = Pt(9)
    r = b.paragraphs[0].add_run(f)
    r.font.size = Pt(8.5)
    r.font.name = 'Consolas'
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
for r in tbl.rows:
    r.cells[0].width, r.cells[1].width = Inches(3.0), Inches(3.5)

doc.save(OUT)
print('saved', OUT, os.path.getsize(OUT) // 1024, 'KB')
